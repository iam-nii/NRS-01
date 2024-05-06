import pandas as pd
import openpyxl

class Save_to_excel:
    def __init__(self, list1: list, list2: list, list3: list, file_name: str, DATA:dict,prod_temp_visc:dict):
        # Define the lists
        x = list1  # coordinates
        y1 = list2  # temperature
        y2 = list3  # viscosity
        data_frame = {
            "Тип материала": DATA["material"],

            "Геометрические параметры канала: ": "",
            "Ширина, м": DATA['width'],
            "Глубина, м": DATA["depth"],
            "Длина ,м": DATA["length"],

            "Параметры свойства материала: ": "",
            "Плотность, кг/м^3": DATA["density"],
            "Удельная теплоемкость, Дж/(кг*С)": DATA["heat_capacity"],
            "Температура плавления, °C": DATA["melting_temperature"],

            "Режимные параметры процесса: ": "",
            "Скорость крышки, м/с": DATA["cover_speed"][0],
            "Температура крышки, °C": DATA['cover_temperature'],

            "Параметры Метода решения управлений модели: ": "",
            "Шаг расчета по длине канала, м": DATA["step"],

            "Эмпирические коэффициенты математической модели: "
            "Коэффициент консистенции материала при температуре приведения, Па*c^n": DATA["consistency_coefficient"],
            "Температурный коэффициент вязкости материала, 1/°C": DATA["temp_viscosity_coefficient"],
            "Температура приведения, °C": DATA["casting_temperature"],
            "Индекс течения материала": DATA["flow_index"],
            "Коэффициент теплоотдачи от крышки канала к материалу, Вт/(м^2*°C)": DATA[
                "cover_heat_transfer_coefficient"],

            "Критериальные показатели процесса: ": "",
            "Производительность": prod_temp_visc["Производительность"],
            "Температура продукта, °C": prod_temp_visc["Температура продукта"],
            "Вязкость продукта": prod_temp_visc["Вязкость продукта"],

            "Показатели экономичности:": "",
            "Память": prod_temp_visc["Память"],
            "Время рассчета": prod_temp_visc["Время рассчета"],
            "Количество операций": prod_temp_visc["Количество операций"],
        }
        labels = [key for key, value in data_frame.items()]
        values = [value for key, value in data_frame.items()]

        print(labels)
        print(values)

        # Create a DataFrame
        temperature_dataframe = pd.DataFrame({'Координаты по длине канала, м': x, 'Температура, °C': y1})
        viscosity_dataframe = pd.DataFrame({'Координаты по длине канала, м': x, 'Вязкость, Па*с': y2})
        report_dataframe = pd.DataFrame({'Входные значения':labels, '':values})

        # Create a new Excel file and add the DataFrame to it
        book = openpyxl.Workbook()
        writer = pd.ExcelWriter(f'{file_name}.xlsx', engine='openpyxl')
        temperature_dataframe.to_excel(writer, sheet_name='Температура')
        viscosity_dataframe.to_excel(writer, sheet_name='Вязкость')


        report_dataframe.to_excel(writer,sheet_name="Отчет")
        writer._save()

        # # Create a chart in the Excel file
        # # Plot temperature graph
        # wb = openpyxl.load_workbook(f'{file_name}.xlsx')
        #
        # temperature_ws = wb['Температура']
        #
        # temperature_chart = openpyxl.chart.ScatterChart()
        # temperature_chart.title = 'График изменения температуры по длине канала'
        # temperature_chart.x_axis.title = 'Координаты по длине канала, м'
        # temperature_chart.y_axis.title = 'Температура, '
        #
        # x_data = openpyxl.chart.Reference(temperature_ws, min_col=2, min_row=2, max_row=len(x) + 1)
        # y_data = openpyxl.chart.Reference(temperature_ws, min_col=3, min_row=2, max_row=len(y1) + 1)
        #
        # temperature_chart.append(openpyxl.chart.Series(y_data, x_data))
        # temperature_ws.add_chart(temperature_chart, "E2")

        # Plot viscosity graph
        # viscosity_ws = wb['Вязкость']
        #
        # viscosity_chart = openpyxl.chart.ScatterChart()
        # viscosity_chart.title = 'График изменения Вязкости по длине канала'
        # viscosity_chart.x_axis.title = 'Координаты по длине канала, м'
        # viscosity_chart.y_axis.title = 'Вязкость'
        #
        # x_data = openpyxl.chart.Reference(viscosity_ws, min_col=2, min_row=2, max_row=len(x) + 1)
        # y_data = openpyxl.chart.Reference(viscosity_ws, min_col=3, min_row=2, max_row=len(y2) + 1)
        #
        # viscosity_chart.append(openpyxl.chart.Series(y_data, x_data))
        # viscosity_ws.add_chart(viscosity_chart, "E2")

        # wb.save(f'{file_name}.xlsx')